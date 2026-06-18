#!/usr/bin/env python3
"""Convert LibreChat mongoexport JSON files to Excel workbooks."""

from __future__ import annotations

import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
EXPORT_DIR = Path(os.environ.get("EXPORT_DIR", ROOT / "export"))
EXCEL_DIR = EXPORT_DIR / "excel"
MAX_CELL = 32700
ILLEGAL_XML = re.compile(r"[\000-\010\013\014\016-\037]")

JSON_LINE = re.compile(r"^\{")


def sanitize_text(text: str) -> str:
    return ILLEGAL_XML.sub("", text)


def parse_mongo_value(value):
    if isinstance(value, dict):
        if "$oid" in value:
            return value["$oid"]
        if "$date" in value:
            raw = value["$date"]
            if isinstance(raw, (int, float)):
                return datetime.fromtimestamp(raw / 1000, tz=timezone.utc).isoformat()
            return str(raw)
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, list):
        return ", ".join(str(parse_mongo_value(item)) for item in value)
    if value is None:
        return ""
    return value


def flatten_doc(doc: dict, fields: list[str]) -> dict:
    row = {}
    for field in fields:
        row[field] = parse_mongo_value(doc.get(field, ""))
    return row


def read_json_lines(path: Path) -> list[dict]:
    docs: list[dict] = []
    skipped = 0
    with path.open(encoding="utf-8", errors="replace") as handle:
        for line_no, line in enumerate(handle, 1):
            line = line.strip()
            if not line or not JSON_LINE.match(line):
                continue
            try:
                docs.append(json.loads(line))
            except json.JSONDecodeError as exc:
                skipped += 1
                print(f"  warning: skipped invalid JSON in {path.name} line {line_no}: {exc.msg}")
    if skipped:
        print(f"  warning: {skipped} invalid line(s) skipped in {path.name}")
    return docs


def autosize_columns(ws) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_len = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            max_len = max(max_len, min(len(str(cell.value)), 80))
        ws.column_dimensions[letter].width = max(10, min(max_len + 2, 60))


def write_sheet(wb: Workbook, title: str, rows: list[dict], headers: list[str]) -> None:
    ws = wb.create_sheet(title=title[:31])
    ws.append(headers)
    for row in rows:
        values = []
        for header in headers:
            value = row.get(header, "")
            text = sanitize_text(str(value) if value is not None else "")
            if len(text) > MAX_CELL:
                text = text[:MAX_CELL] + "…"
            values.append(text)
        ws.append(values)
    ws.freeze_panes = "A2"
    autosize_columns(ws)


def build_users(docs: list[dict]) -> tuple[list[dict], list[str]]:
    headers = [
        "_id",
        "name",
        "username",
        "email",
        "emailVerified",
        "role",
        "provider",
        "createdAt",
        "updatedAt",
    ]
    rows = [flatten_doc(doc, headers) for doc in docs]
    return rows, headers


def build_conversations(docs: list[dict]) -> tuple[list[dict], list[str]]:
    headers = [
        "_id",
        "conversationId",
        "user",
        "title",
        "endpoint",
        "endpointType",
        "model",
        "isArchived",
        "messageCount",
        "createdAt",
        "updatedAt",
    ]
    rows = []
    for doc in docs:
        row = flatten_doc(doc, headers)
        messages = doc.get("messages", [])
        row["messageCount"] = len(messages) if isinstance(messages, list) else 0
        rows.append(row)
    return rows, headers


def build_messages(docs: list[dict]) -> tuple[list[dict], list[str]]:
    headers = [
        "_id",
        "messageId",
        "conversationId",
        "user",
        "sender",
        "text",
        "isCreatedByUser",
        "endpoint",
        "model",
        "tokenCount",
        "error",
        "createdAt",
        "updatedAt",
    ]
    rows = [flatten_doc(doc, headers) for doc in docs]
    return rows, headers


def build_combined_rows(
    users: list[dict], conversations: list[dict], messages: list[dict]
) -> tuple[list[dict], list[str]]:
    users_by_id = {parse_mongo_value(doc.get("_id")): doc for doc in users}
    conv_by_id = {doc.get("conversationId", ""): doc for doc in conversations}

    headers = [
        "user_id",
        "user_name",
        "user_email",
        "user_role",
        "conversation_id",
        "conversation_title",
        "conversation_endpoint",
        "conversation_model",
        "conversation_created_at",
        "message_id",
        "sender",
        "message_text",
        "is_user_message",
        "token_count",
        "message_created_at",
    ]

    rows: list[dict] = []
    for message in messages:
        user_id = parse_mongo_value(message.get("user", ""))
        user = users_by_id.get(user_id, {})
        conv_id = message.get("conversationId", "")
        conv = conv_by_id.get(conv_id, {})

        rows.append(
            {
                "user_id": user_id,
                "user_name": parse_mongo_value(user.get("name", "")),
                "user_email": parse_mongo_value(user.get("email", "")),
                "user_role": parse_mongo_value(user.get("role", "")),
                "conversation_id": conv_id,
                "conversation_title": parse_mongo_value(conv.get("title", "")),
                "conversation_endpoint": parse_mongo_value(conv.get("endpoint", "")),
                "conversation_model": parse_mongo_value(conv.get("model", "")),
                "conversation_created_at": parse_mongo_value(conv.get("createdAt", "")),
                "message_id": parse_mongo_value(message.get("messageId", message.get("_id", ""))),
                "sender": parse_mongo_value(message.get("sender", "")),
                "message_text": parse_mongo_value(message.get("text", "")),
                "is_user_message": parse_mongo_value(message.get("isCreatedByUser", "")),
                "token_count": parse_mongo_value(message.get("tokenCount", "")),
                "message_created_at": parse_mongo_value(message.get("createdAt", "")),
            }
        )

    rows.sort(
        key=lambda row: (
            str(row.get("conversation_created_at", "")),
            str(row.get("message_created_at", "")),
        )
    )
    return rows, headers


def load_combined_data(export_dir: Path | None = None) -> tuple[list[dict], list[str], dict[str, int]]:
    """Load joined user/conversation/message rows from JSON export files."""
    base = export_dir or EXPORT_DIR
    users = read_json_lines(base / "users.json")
    conversations = read_json_lines(base / "conversations.json")
    messages = read_json_lines(base / "messages.json")
    rows, headers = build_combined_rows(users, conversations, messages)
    stats = {
        "users": len(users),
        "conversations": len(conversations),
        "messages": len(messages),
        "rows": len(rows),
    }
    return rows, headers, stats


def main() -> None:
    EXCEL_DIR.mkdir(parents=True, exist_ok=True)

    users_path = EXPORT_DIR / "users.json"
    conversations_path = EXPORT_DIR / "conversations.json"
    messages_path = EXPORT_DIR / "messages.json"

    print("Reading JSON exports...")
    users = read_json_lines(users_path)
    conversations = read_json_lines(conversations_path)
    messages = read_json_lines(messages_path)
    print(f"  users: {len(users)}")
    print(f"  conversations: {len(conversations)}")
    print(f"  messages: {len(messages)}")

    combined_rows, combined_headers = build_combined_rows(users, conversations, messages)

    combined_path = EXCEL_DIR / "librechat-combined.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "combined"
    ws.append(combined_headers)
    for row in combined_rows:
        values = []
        for header in combined_headers:
            text = sanitize_text(str(row.get(header, "")))
            if len(text) > MAX_CELL:
                text = text[:MAX_CELL] + "…"
            values.append(text)
        ws.append(values)
    ws.freeze_panes = "A2"
    autosize_columns(ws)
    wb.save(combined_path)
    print(f"Saved: {combined_path} ({len(combined_rows)} rows)")


if __name__ == "__main__":
    main()
