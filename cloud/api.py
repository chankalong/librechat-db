"""HTTP API to fetch LibreChat export files for dashboards (Power BI, etc.)."""

from __future__ import annotations

import csv
import io
import json
import os
import secrets
from typing import Annotated

from fastapi import Depends, FastAPI, Header, HTTPException, Query
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse

from storage import ExportArtifact, list_available_dates, resolve_export

app = FastAPI(
    title="LibreChat Export API",
    description="Fetch daily Excel/CSV/JSON exports for external dashboards.",
    version="1.0.0",
)

API_KEY = os.environ.get("API_KEY", "")
EXCEL_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _load_json_to_excel():
    import importlib.util

    spec = importlib.util.spec_from_file_location("json_to_excel", "/app/json-to-excel.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("json-to-excel.py not found")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def verify_api_key(
    authorization: Annotated[str | None, Header()] = None,
    x_api_key: Annotated[str | None, Header(alias="X-API-Key")] = None,
) -> None:
    if not API_KEY:
        raise HTTPException(
            status_code=503,
            detail="API_KEY is not configured on the server",
        )

    token = None
    if x_api_key:
        token = x_api_key.strip()
    elif authorization and authorization.lower().startswith("bearer "):
        token = authorization[7:].strip()

    if not token or not secrets.compare_digest(token, API_KEY):
        raise HTTPException(status_code=401, detail="Invalid or missing API key")


def _artifact_or_404(export_date: str | None = None) -> ExportArtifact:
    try:
        return resolve_export(export_date)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


def _metadata_payload(artifact: ExportArtifact) -> dict:
    stat = artifact.excel_path.stat()
    payload = {
        "export_date": artifact.export_date,
        "source": artifact.source,
        "file_name": artifact.excel_path.name,
        "file_size_bytes": stat.st_size,
        "updated_at": stat.st_mtime,
        "available_dates": list_available_dates(),
        "endpoints": {
            "excel": "/api/v1/export/latest.xlsx",
            "csv": "/api/v1/export/latest.csv",
            "json": "/api/v1/export/latest.json",
            "metadata": "/api/v1/export/metadata",
        },
    }
    if artifact.export_dir:
        payload["json_export_dir"] = str(artifact.export_dir)
    return payload


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/api/v1/export/metadata")
def export_metadata(_: Annotated[None, Depends(verify_api_key)] = None) -> dict:
    artifact = _artifact_or_404()
    return _metadata_payload(artifact)


@app.get("/api/v1/export/dates")
def export_dates(_: Annotated[None, Depends(verify_api_key)] = None) -> dict:
    dates = list_available_dates()
    return {"dates": dates, "count": len(dates)}


@app.get("/api/v1/export/latest.xlsx")
def download_latest_excel(_: Annotated[None, Depends(verify_api_key)] = None) -> FileResponse:
    artifact = _artifact_or_404()
    filename = f"librechat-combined-{artifact.export_date}.xlsx"
    return FileResponse(
        path=artifact.excel_path,
        media_type=EXCEL_MEDIA,
        filename=filename,
    )


@app.get("/api/v1/export/{export_date}.xlsx")
def download_dated_excel(
    export_date: str,
    _: Annotated[None, Depends(verify_api_key)] = None,
) -> FileResponse:
    artifact = _artifact_or_404(export_date)
    return FileResponse(
        path=artifact.excel_path,
        media_type=EXCEL_MEDIA,
        filename=artifact.excel_path.name,
    )


@app.get("/api/v1/export/latest.csv")
def download_latest_csv(_: Annotated[None, Depends(verify_api_key)] = None) -> StreamingResponse:
    artifact = _artifact_or_404()
    rows, headers = _load_rows(artifact)
    return _csv_response(rows, headers, f"librechat-combined-{artifact.export_date}.csv")


@app.get("/api/v1/export/{export_date}.csv")
def download_dated_csv(
    export_date: str,
    _: Annotated[None, Depends(verify_api_key)] = None,
) -> StreamingResponse:
    artifact = _artifact_or_404(export_date)
    rows, headers = _load_rows(artifact)
    return _csv_response(rows, headers, f"librechat-combined-{export_date}.csv")


@app.get("/api/v1/export/latest.json")
def download_latest_json(
    page: Annotated[int, Query(ge=1)] = 1,
    limit: Annotated[int, Query(ge=1, le=5000)] = 1000,
    _: Annotated[None, Depends(verify_api_key)] = None,
) -> JSONResponse:
    artifact = _artifact_or_404()
    rows, headers, stats = _load_rows_with_stats(artifact)
    start = (page - 1) * limit
    end = start + limit
    page_rows = rows[start:end]
    return JSONResponse(
        {
            "export_date": artifact.export_date,
            "page": page,
            "limit": limit,
            "total_rows": len(rows),
            "total_pages": max(1, (len(rows) + limit - 1) // limit),
            "stats": stats,
            "headers": headers,
            "rows": page_rows,
        }
    )


def _load_rows(artifact: ExportArtifact) -> tuple[list[dict], list[str]]:
    rows, headers, _ = _load_rows_with_stats(artifact)
    return rows, headers


def _load_rows_with_stats(artifact: ExportArtifact) -> tuple[list[dict], list[str], dict[str, int]]:
    module = _load_json_to_excel()
    if artifact.export_dir:
        return module.load_combined_data(artifact.export_dir)

    # Fall back to reading rows from the Excel file when only S3 xlsx is available.
    from openpyxl import load_workbook

    workbook = load_workbook(artifact.excel_path, read_only=True, data_only=True)
    sheet = workbook.active
    row_iter = sheet.iter_rows(values_only=True)
    headers = [str(cell) if cell is not None else "" for cell in next(row_iter)]
    rows: list[dict] = []
    for values in row_iter:
        rows.append(
            {
                headers[i]: "" if values[i] is None else str(values[i])
                for i in range(len(headers))
            }
        )
    workbook.close()
    stats = {"users": 0, "conversations": 0, "messages": len(rows), "rows": len(rows)}
    return rows, headers, stats


def _csv_response(rows: list[dict], headers: list[str], filename: str) -> StreamingResponse:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    buffer.seek(0)
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
